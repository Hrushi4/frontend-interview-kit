#!/usr/bin/env python3
"""Convert the Markdown study files into .docx files that upload cleanly to Google Docs.

Usage:  python3 build_docx.py            # converts every frontend-interview-kit/*.md
        python3 build_docx.py 05         # converts only files starting with 05

Requires: pip install python-docx markdown-it-py
Output:   google-docs/<same-name>.docx
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from markdown_it import MarkdownIt

ROOT = Path(__file__).parent
SRC = ROOT / "frontend-interview-kit"
OUT = ROOT / "google-docs"

BODY_FONT = "Arial"
CODE_FONT = "Courier New"
CODE_BG = "F3F4F6"
QUOTE_BG = "EEF4FF"
HEADER_BG = "1F3A5F"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)

QUESTION_RE = re.compile(r"^\*\*Q\d+\.")

md = MarkdownIt("commonmark").enable("table").enable("strikethrough")


# Elements that must come after <w:pBdr>/<w:shd> inside <w:pPr> (OOXML order).
PPR_AFTER_SHD = (
    "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
    "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
    "w:sectPr", "w:pPrChange",
)
TCPR_AFTER_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
                  "w:vAlign", "w:hideMark")


def shade(cell_or_par, hex_color):
    """Apply a background fill to a table cell or paragraph."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    if hasattr(cell_or_par, "_tc"):
        cell_or_par._tc.get_or_add_tcPr().insert_element_before(shd, *TCPR_AFTER_SHD)
    else:
        cell_or_par._p.get_or_add_pPr().insert_element_before(shd, *PPR_AFTER_SHD)


def left_border(par, hex_color, size=18):
    p_pr = par._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), hex_color)
    borders.append(left)
    p_pr.insert_element_before(borders, "w:shd", *PPR_AFTER_SHD)


def bottom_rule(par):
    p_pr = par._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFC7D5")
    borders.append(bottom)
    p_pr.insert_element_before(borders, "w:shd", *PPR_AFTER_SHD)


def set_cell_margins(table, cm=0.15):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(cm * 567)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.insert_element_before(mar, "w:tblLook", "w:tblCaption", "w:tblDescription")


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    sizes = {1: 22, 2: 16, 3: 13, 4: 12}
    for level, size in sizes.items():
        st = styles[f"Heading {level}"]
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = ACCENT
        rpr = st.element.get_or_add_rPr()
        rpr.rFonts.set(qn("w:ascii"), BODY_FONT)
        rpr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        st.paragraph_format.space_before = Pt(18 if level <= 2 else 12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin = sec.bottom_margin = Cm(2.0)


def add_runs(par, children, base=None):
    """Render inline markdown tokens as runs on a paragraph."""
    state = dict(bold=False, italic=False, strike=False, link=None)
    if base:
        state.update(base)
    for tok in children or []:
        t = tok.type
        if t == "text":
            run = par.add_run(tok.content)
            run.bold = state["bold"] or None
            run.italic = state["italic"] or None
            run.font.strike = state["strike"] or None
            if state["link"]:
                run.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)
                run.underline = True
        elif t == "code_inline":
            run = par.add_run(tok.content)
            run.font.name = CODE_FONT
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0xB4, 0x23, 0x18)
            run.bold = state["bold"] or None
        elif t == "strong_open":
            state["bold"] = True
        elif t == "strong_close":
            state["bold"] = False
        elif t == "em_open":
            state["italic"] = True
        elif t == "em_close":
            state["italic"] = False
        elif t == "s_open":
            state["strike"] = True
        elif t == "s_close":
            state["strike"] = False
        elif t == "link_open":
            state["link"] = tok.attrs.get("href")
        elif t == "link_close":
            href = state["link"]
            state["link"] = None
            if href and href.startswith("http"):
                par.add_run(f" ({href})").font.size = Pt(9)
        elif t in ("softbreak", "hardbreak"):
            par.add_run().add_break()
        elif t == "html_inline":
            if tok.content.lower().startswith("<br"):
                par.add_run().add_break()
            else:
                par.add_run(tok.content)
        elif t == "image":
            par.add_run(f"[image: {tok.content}]")


def add_code_block(doc, code, lang):
    lines = code.rstrip("\n").split("\n")
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    shade(cell, CODE_BG)
    par = cell.paragraphs[0]
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.line_spacing = 1.0
    if lang:
        label = par.add_run(lang.upper() + "\n")
        label.font.size = Pt(8)
        label.font.bold = True
        label.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        label.font.name = BODY_FONT
    for i, line in enumerate(lines):
        run = par.add_run(line.replace("\t", "    "))
        run.font.name = CODE_FONT
        run.element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)
        run.font.size = Pt(9.5)
        if i < len(lines) - 1:
            run.add_break()
    set_cell_margins(table, 0.25)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_table(doc, tokens, i):
    """Consume tokens from table_open to table_close and render a table."""
    rows, current, is_header = [], None, []
    j = i + 1
    while tokens[j].type != "table_close":
        tok = tokens[j]
        if tok.type == "tr_open":
            current = []
        elif tok.type == "tr_close":
            rows.append(current)
        elif tok.type in ("th_open", "td_open"):
            inline = tokens[j + 1]
            current.append((tok.type == "th_open", inline.children))
        j += 1
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(rows):
        for c_idx, (head, children) in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            add_runs(par, children, {"bold": head})
            for run in par.runs:
                run.font.size = Pt(10)
                if head:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if head:
                shade(cell, HEADER_BG)
            elif r_idx % 2 == 0:
                shade(cell, "F7F9FC")
    set_cell_margins(table, 0.12)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return j


def convert(md_path: Path, out_path: Path):
    text = md_path.read_text(encoding="utf-8")
    tokens = md.parse(text)
    doc = Document()
    setup_styles(doc)
    core = doc.core_properties
    first = text.splitlines()[0].lstrip("# ").strip()
    core.title = first
    core.author = "Frontend Interview Kit"

    list_stack = []  # each entry: [ordered, counter]
    quote_depth = 0
    i = 0
    while i < len(tokens):
        tok = tokens[i]
        t = tok.type
        if t == "heading_open":
            level = int(tok.tag[1])
            inline = tokens[i + 1]
            par = doc.add_heading(level=min(level, 4))
            add_runs(par, inline.children)
            if level == 2:
                bottom_rule(par)
            i += 3
            continue
        if t == "paragraph_open":
            inline = tokens[i + 1]
            if list_stack:
                ordered, _ = list_stack[-1]
                depth = len(list_stack)
                if ordered:
                    list_stack[-1][1] += 1
                    par = doc.add_paragraph()
                    par.paragraph_format.left_indent = Cm(0.75 * depth)
                    par.paragraph_format.first_line_indent = Cm(-0.6)
                    par.add_run(f"{list_stack[-1][1]}. ")
                else:
                    par = doc.add_paragraph(style="List Bullet 2" if depth > 1 else "List Bullet")
                par.paragraph_format.space_after = Pt(2)
            else:
                par = doc.add_paragraph()
            if quote_depth:
                shade(par, QUOTE_BG)
                left_border(par, "3B82F6")
                par.paragraph_format.left_indent = Cm(0.3)
            add_runs(par, inline.children)
            if not list_stack and QUESTION_RE.match(inline.content):
                # "**Q12. …**" lines: make questions easy to scan in the doc
                par.paragraph_format.space_before = Pt(10)
                par.paragraph_format.keep_with_next = True
                for run in par.runs:
                    run.font.size = Pt(12)
                    run.font.color.rgb = ACCENT
            i += 3
            continue
        if t == "bullet_list_open":
            list_stack.append([False, 0])
        elif t == "ordered_list_open":
            list_stack.append([True, int(tok.attrs.get("start", 1)) - 1])
        elif t in ("bullet_list_close", "ordered_list_close"):
            list_stack.pop()
            if not list_stack:
                doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif t == "blockquote_open":
            quote_depth += 1
        elif t == "blockquote_close":
            quote_depth -= 1
        elif t in ("fence", "code_block"):
            add_code_block(doc, tok.content, (tok.info or "").strip())
        elif t == "table_open":
            i = build_table(doc, tokens, i)
        elif t == "hr":
            par = doc.add_paragraph()
            bottom_rule(par)
        elif t == "html_block":
            pass
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def main():
    prefix = sys.argv[1] if len(sys.argv) > 1 else ""
    files = sorted(p for p in SRC.glob("*.md") if p.name.startswith(prefix))
    for path in files:
        out = OUT / (path.stem + ".docx")
        convert(path, out)
        print(f"built {out.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
